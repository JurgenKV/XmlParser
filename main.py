import os
import html
import threading
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from collections import Counter, defaultdict

from lxml import etree as ET

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont


# ============================================================
# 1. АНАЛИЗ СТРУКТУРЫ XML
# ============================================================

def analyze_structure(path, sample_limit=200000):
    """
    Потоково собирает статистику по тегам:
      - сколько раз встречается каждый тег,
      - какие теги-дети у него бывают (Counter).
    Безопасно обрабатывает комментарии и processing instructions.
    """
    stats = defaultdict(lambda: {"count": 0, "children": Counter()})

    context = ET.iterparse(path, events=("start", "end"))
    _, root = next(context)

    stack = []
    processed = 0

    for event, elem in context:
        tag = elem.tag if isinstance(elem.tag, str) else None

        if event == "start":
            stack.append(elem)
            if tag is not None:
                stats[tag]["count"] += 1
        else:  # end
            if stack:
                stack.pop()

            if tag is not None and stack:
                parent = stack[-1]
                ptag = parent.tag if isinstance(parent.tag, str) else None
                if ptag is not None:
                    stats[ptag]["children"][tag] += 1

            # освобождаем память
            elem.clear()
            while elem.getprevious() is not None:
                del elem.getparent()[0]

            processed += 1
            if processed >= sample_limit:
                break

    del root
    return stats


def detect_record_tag(stats):
    """
    Эвристика: 'запись' — тег, который встречается много раз
    и имеет стабильный набор детей.
    """
    best = None
    best_score = 0

    for tag, info in stats.items():
        cnt = info["count"]
        n_children = len(info["children"])
        if cnt < 2 or n_children < 1:
            continue

        total_child_uses = sum(info["children"].values())
        score = cnt * min(n_children, 5) + total_child_uses

        # если у всех записей ровно один и тот же дочерний тег —
        # это, скорее всего, не запись, а обёртка
        if len(info["children"]) == 1:
            score *= 0.6

        if score > best_score:
            best_score = score
            best = tag

    return best


# ============================================================
# 2. ИЗВЛЕЧЕНИЕ ЗАПИСЕЙ
# ============================================================

def _path_under(root, node):
    """Путь тега node относительно root (без самого root)."""
    parts = []
    cur = node
    guard = 0
    while cur is not None and cur is not root:
        parts.append(cur.tag if isinstance(cur.tag, str) else "?")
        cur = cur.getparent()
        guard += 1
        if guard > 1000:  # защита от бесконечного цикла
            break
    return ".".join(reversed(parts)) or (
        node.tag if isinstance(node.tag, str) else "?")


def extract_records(path, record_tag, progress_cb=None, cancel_flag=None,
                    max_records=None):
    """
    Потоково извлекает записи по тегу record_tag.
    Возвращает генератор dict: {"field": value, "@attr": value, ...}.
    """
    context = ET.iterparse(path, events=("end",), tag=record_tag)

    count = 0
    for _, elem in context:
        if cancel_flag and cancel_flag.is_set():
            return

        rec = {}

        # атрибуты самого элемента
        for k, v in elem.attrib.items():
            rec[f"@{k}"] = v

        # вложенные поля: ключ — путь тега от записи
        for child in elem.iter():
            if child is elem:
                continue
            if not isinstance(child.tag, str):
                continue
            text = (child.text or "").strip()
            if text:
                key = _path_under(elem, child)
                if key in rec:
                    i = 2
                    while f"{key}_{i}" in rec:
                        i += 1
                    key = f"{key}_{i}"
                rec[key] = text

        # прямой текст внутри самой записи
        direct_text = (elem.text or "").strip()
        if direct_text:
            rec["_text"] = direct_text

        yield rec

        # освобождаем память
        elem.clear()
        while elem.getprevious() is not None:
            del elem.getparent()[0]

        count += 1
        if progress_cb and count % 1000 == 0:
            progress_cb(count)
        if max_records and count >= max_records:
            return


# ============================================================
# 3. DOCX ЭКСПОРТ
# ============================================================

def _add_table_docx(doc, records, columns, title):
    doc.add_heading(title, level=1)

    table = doc.add_table(rows=1, cols=len(columns))
    try:
        table.style = "Light Grid Accent 1"
    except KeyError:
        # если стиля нет в шаблоне — оставим дефолтный
        pass
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr = table.rows[0].cells
    for i, col in enumerate(columns):
        hdr[i].text = str(col)
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True

    for rec in records:
        row = table.add_row().cells
        for i, col in enumerate(columns):
            row[i].text = str(rec.get(col, ""))

    return table


def _add_cards_docx(doc, records, columns, title):
    doc.add_heading(title, level=1)
    for i, rec in enumerate(records, 1):
        p = doc.add_paragraph()
        run = p.add_run(f"Запись {i}")
        run.bold = True
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0x00, 0x4A, 0x99)

        for col in columns:
            if col not in rec:
                continue
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            r1 = p.add_run(f"{col}: ")
            r1.bold = True
            p.add_run(str(rec[col]))
        doc.add_paragraph()


def export_docx(xml_path, out_path, record_tag,
                limit=100000, progress_cb=None, cancel_flag=None,
                as_cards=False):
    doc = Document()

    title = doc.add_heading("Отчёт из XML", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(f"Источник: {os.path.basename(xml_path)}")
    r.italic = True

    records = list(extract_records(xml_path, record_tag,
                                   progress_cb, cancel_flag,
                                   max_records=limit))

    if not records:
        doc.add_paragraph("Записи не найдены.")
        doc.save(out_path)
        return

    # колонки — по порядку первого появления
    seen = {}
    for rec in records:
        for k in rec:
            if k not in seen:
                seen[k] = True
    columns = list(seen.keys())

    if as_cards:
        _add_cards_docx(doc, records, columns, "Записи")
    else:
        _add_table_docx(doc, records, columns, "Таблица записей")

    doc.save(out_path)


# ============================================================
# 4. HTML ЭКСПОРТ
# ============================================================

HTML_HEAD = """<!DOCTYPE html>
<html lang="ru"><head><meta charset="utf-8">
<title>{title}</title>
<style>
  body {{ font-family: -apple-system, Segoe UI, Arial, sans-serif;
         margin: 2em; color: #222; }}
  h1 {{ border-bottom: 2px solid #444; padding-bottom: .3em; }}
  table {{ border-collapse: collapse; margin: 1em 0; width: 100%; }}
  th, td {{ border: 1px solid #ccc; padding: 6px 10px; text-align: left;
            vertical-align: top; }}
  th {{ background: #f0f4f8; }}
  tr:nth-child(even) {{ background: #fafbfc; }}
  .card {{ border: 1px solid #ddd; border-radius: 6px; padding: 12px 16px;
           margin: 8px 0; background: #fafbfc; }}
  .card h3 {{ margin-top: 0; color: #004a99; }}
  .field {{ margin: 2px 0; }}
  .field b {{ color: #333; }}
</style></head><body>
<h1>{title}</h1>
<p><i>Источник: {source}<br>Записей: {count}</i></p>
"""


def _esc(s):
    return html.escape(str(s))


def export_html(xml_path, out_path, record_tag,
                limit=100000, progress_cb=None, cancel_flag=None,
                as_cards=False):
    records = list(extract_records(xml_path, record_tag,
                                   progress_cb, cancel_flag,
                                   max_records=limit))

    seen = {}
    for rec in records:
        for k in rec:
            if k not in seen:
                seen[k] = True
    columns = list(seen.keys())

    with open(out_path, "w", encoding="utf-8") as f:
        f.write(HTML_HEAD.format(
            title="Отчёт из XML",
            source=_esc(os.path.basename(xml_path)),
            count=len(records),
        ))

        if as_cards:
            for i, rec in enumerate(records, 1):
                f.write(f'<div class="card"><h3>Запись {i}</h3>\n')
                for col in columns:
                    if col in rec:
                        f.write(f'<div class="field"><b>{_esc(col)}:</b> '
                                f'{_esc(rec[col])}</div>\n')
                f.write("</div>\n")
        else:
            f.write("<table><thead><tr>")
            for col in columns:
                f.write(f"<th>{_esc(col)}</th>")
            f.write("</tr></thead><tbody>\n")
            for rec in records:
                f.write("<tr>")
                for col in columns:
                    f.write(f"<td>{_esc(rec.get(col, ''))}</td>")
                f.write("</tr>\n")
            f.write("</tbody></table>\n")

        f.write("</body></html>\n")


# ============================================================
# 5. PDF ЭКСПОРТ
# ============================================================

def register_cyrillic_font():
    for path in [
        "C:/Windows/Fonts/DejaVuSans.ttf",
        "C:/Windows/Fonts/arial.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/Library/Fonts/Arial Unicode.ttf",
    ]:
        if os.path.exists(path):
            try:
                pdfmetrics.registerFont(TTFont("UnicodeFont", path))
                return "UnicodeFont"
            except Exception:
                continue
    return "Helvetica"


def export_pdf(xml_path, out_path, record_tag,
               limit=100000, progress_cb=None, cancel_flag=None,
               as_cards=False):
    font = register_cyrillic_font()

    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("H1", parent=styles["Heading1"], fontName=font,
                        fontSize=16, leading=20)
    normal = ParagraphStyle("N", parent=styles["Normal"], fontName=font,
                            fontSize=9, leading=12)
    small = ParagraphStyle("S", parent=styles["Normal"], fontName=font,
                           fontSize=8, leading=10, textColor=colors.grey)

    records = list(extract_records(xml_path, record_tag,
                                   progress_cb, cancel_flag,
                                   max_records=limit))

    seen = {}
    for rec in records:
        for k in rec:
            if k not in seen:
                seen[k] = True
    columns = list(seen.keys())

    use_landscape = len(columns) > 5
    pagesize = landscape(A4) if use_landscape else A4

    doc = SimpleDocTemplate(
        out_path, pagesize=pagesize,
        leftMargin=1.5 * cm, rightMargin=1.5 * cm,
        topMargin=1.5 * cm, bottomMargin=1.5 * cm,
    )

    story = []
    story.append(Paragraph("Отчёт из XML", h1))
    story.append(Paragraph(
        f"Источник: {_esc(os.path.basename(xml_path))} &nbsp;|&nbsp; "
        f"Записей: {len(records)}", small))
    story.append(Spacer(1, 8))

    if as_cards:
        for i, rec in enumerate(records, 1):
            story.append(Paragraph(f"<b>Запись {i}</b>", normal))
            for col in columns:
                if col in rec:
                    story.append(Paragraph(
                        f"<b>{_esc(col)}:</b> {_esc(rec[col])}", normal))
            story.append(Spacer(1, 6))
    else:
        data = [[Paragraph(f"<b>{_esc(c)}</b>", normal) for c in columns]]
        for rec in records:
            data.append([
                Paragraph(_esc(rec.get(c, "")), normal) for c in columns
            ])

        avail = (pagesize[0] - 3 * cm) / max(len(columns), 1)
        col_widths = [avail] * len(columns)

        tbl = Table(data, colWidths=col_widths, repeatRows=1)
        tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f0f4f8")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        story.append(tbl)

    doc.build(story)


# ============================================================
# 6. GUI
# ============================================================

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("XML → отчёт (DOCX / PDF / HTML)")
        self.geometry("720x580")
        self.resizable(False, False)

        self.xml_path = tk.StringVar()
        self.format_var = tk.StringVar(value="html")
        self.record_tag = tk.StringVar(value="")
        self.limit = tk.IntVar(value=100000)
        self.view_mode = tk.StringVar(value="table")

        self.cancel_flag = threading.Event()
        self.detected_tags = []

        self._build_ui()

    def _build_ui(self):
        pad = {"padx": 10, "pady": 6}

        # --- файл ---
        frame_file = ttk.LabelFrame(self, text="1. XML-файл")
        frame_file.pack(fill="x", **pad)
        ttk.Entry(frame_file, textvariable=self.xml_path, width=65).pack(
            side="left", padx=6, pady=6, fill="x", expand=True)
        ttk.Button(frame_file, text="Обзор…", command=self.choose_file).pack(
            side="right", padx=6, pady=6)
        ttk.Button(frame_file, text="Анализ структуры",
                   command=self.run_analysis).pack(side="right", padx=6, pady=6)

        # --- запись ---
        frame_rec = ttk.LabelFrame(
            self, text="2. Тег записи (пусто = автоопределение)")
        frame_rec.pack(fill="x", **pad)
        self.cmb_tag = ttk.Combobox(frame_rec, textvariable=self.record_tag,
                                    values=[], width=40)
        self.cmb_tag.pack(side="left", padx=6, pady=6, fill="x", expand=True)
        ttk.Label(frame_rec, text="Макс. записей:").pack(side="left", padx=6)
        ttk.Entry(frame_rec, textvariable=self.limit, width=10).pack(
            side="left", padx=6)

        # --- представление ---
        frame_view = ttk.LabelFrame(self, text="3. Представление")
        frame_view.pack(fill="x", **pad)
        ttk.Radiobutton(frame_view, text="Таблица",
                        value="table", variable=self.view_mode).pack(
            side="left", padx=12, pady=6)
        ttk.Radiobutton(frame_view, text="Карточки",
                        value="cards", variable=self.view_mode).pack(
            side="left", padx=12, pady=6)

        # --- формат ---
        frame_fmt = ttk.LabelFrame(self, text="4. Формат вывода")
        frame_fmt.pack(fill="x", **pad)
        for fmt, label in [("html", "HTML"), ("docx", "Word (.docx)"),
                           ("pdf", "PDF")]:
            ttk.Radiobutton(frame_fmt, text=label, value=fmt,
                            variable=self.format_var).pack(
                side="left", padx=12, pady=6)

        # --- прогресс ---
        frame_prog = ttk.LabelFrame(self, text="Прогресс")
        frame_prog.pack(fill="x", **pad)
        self.progress = ttk.Progressbar(frame_prog, mode="indeterminate")
        self.progress.pack(fill="x", padx=6, pady=8)

        # --- кнопки ---
        frame_btn = ttk.Frame(self)
        frame_btn.pack(fill="x", **pad)
        self.btn_convert = ttk.Button(frame_btn, text="Сформировать отчёт",
                                      command=self.convert)
        self.btn_convert.pack(side="left", padx=6)
        self.btn_cancel = ttk.Button(frame_btn, text="Отмена",
                                     command=self.cancel, state="disabled")
        self.btn_cancel.pack(side="left", padx=6)

        self.status = ttk.Label(self, text="Готов к работе", foreground="gray")
        self.status.pack(pady=4)

    # ---------- действия ----------

    def choose_file(self):
        path = filedialog.askopenfilename(
            title="Выберите XML-файл",
            filetypes=[("XML files", "*.xml"), ("All files", "*.*")])
        if path:
            self.xml_path.set(path)

    def cancel(self):
        self.cancel_flag.set()

    def _set_status(self, text, color="gray"):
        self.after(0, lambda: self.status.config(text=text, foreground=color))

    def _progress_cb(self, count):
        if count % 5000 == 0:
            self._set_status(f"Обработано записей: {count:,}", "blue")

    def run_analysis(self):
        xml_file = self.xml_path.get().strip()
        if not xml_file or not os.path.isfile(xml_file):
            messagebox.showwarning("Внимание", "Сначала выберите XML-файл")
            return

        self.progress.start(10)
        self._set_status("Анализ структуры…", "blue")
        self.btn_convert.config(state="disabled")

        def worker():
            try:
                stats = analyze_structure(xml_file, sample_limit=200000)
                guessed = detect_record_tag(stats)

                candidates = sorted(
                    ((tag, info["count"], len(info["children"]))
                     for tag, info in stats.items()
                     if info["count"] >= 2 and len(info["children"]) >= 1),
                    key=lambda x: -x[1]
                )
                top = [c[0] for c in candidates[:20]]

                def apply():
                    self.detected_tags = top
                    self.cmb_tag["values"] = top
                    if guessed and not self.record_tag.get():
                        self.record_tag.set(guessed)
                    self._set_status(
                        f"Кандидатов: {len(top)}. Предполагаемая запись: "
                        f"{guessed or '—'}", "green")
                self.after(0, apply)

            except Exception as exc:
                err = f"{type(exc).__name__}: {exc}"
                self.after(0, self._show_error, "Ошибка анализа", err)
            finally:
                self.after(0, self._reset_ui)

        threading.Thread(target=worker, daemon=True).start()

    def convert(self):
        xml_file = self.xml_path.get().strip()
        if not xml_file or not os.path.isfile(xml_file):
            messagebox.showwarning("Внимание", "Выберите существующий XML-файл")
            return

        fmt = self.format_var.get()
        ext = fmt
        record_tag = self.record_tag.get().strip()
        as_cards = self.view_mode.get() == "cards"
        limit = max(1, int(self.limit.get() or 100000))

        out_path = filedialog.asksaveasfilename(
            title="Сохранить как",
            defaultextension=f".{ext}",
            filetypes=[(f"{ext.upper()} files", f"*.{ext}")],
            initialfile=os.path.splitext(os.path.basename(xml_file))[0]
                        + f"_report.{ext}")
        if not out_path:
            return

        self.cancel_flag.clear()
        self.btn_convert.config(state="disabled")
        self.btn_cancel.config(state="normal")
        self.progress.start(10)
        self._set_status("Формирование отчёта…", "blue")

        def worker():
            try:
                tag = record_tag
                if not tag:
                    self._set_status("Определение тега записи…", "blue")
                    stats = analyze_structure(xml_file, sample_limit=200000)
                    tag = detect_record_tag(stats)
                    if not tag:
                        raise RuntimeError(
                            "Не удалось автоматически определить тег записи. "
                            "Укажите его вручную в поле «Тег записи».")
                    self.after(0, lambda t=tag: self.record_tag.set(t))

                if self.cancel_flag.is_set():
                    raise RuntimeError("Отменено")

                if fmt == "html":
                    export_html(xml_file, out_path, tag,
                                limit=limit,
                                progress_cb=self._progress_cb,
                                cancel_flag=self.cancel_flag,
                                as_cards=as_cards)
                elif fmt == "docx":
                    export_docx(xml_file, out_path, tag,
                                limit=limit,
                                progress_cb=self._progress_cb,
                                cancel_flag=self.cancel_flag,
                                as_cards=as_cards)
                elif fmt == "pdf":
                    export_pdf(xml_file, out_path, tag,
                               limit=limit,
                               progress_cb=self._progress_cb,
                               cancel_flag=self.cancel_flag,
                               as_cards=as_cards)
                else:
                    raise ValueError(f"Неизвестный формат: {fmt}")

                self.after(0, self._show_success, f"Файл сохранён:\n{out_path}")

            except Exception as exc:
                err = f"{type(exc).__name__}: {exc}"
                self.after(0, self._show_error, "Ошибка", err)
            finally:
                self.after(0, self._reset_ui)

        threading.Thread(target=worker, daemon=True).start()

    def _show_error(self, title, message):
        self.status.config(text=title, foreground="red")
        messagebox.showerror(title, message)

    def _show_success(self, message):
        self.status.config(text="Готово", foreground="green")
        messagebox.showinfo("Успех", message)

    def _reset_ui(self):
        self.progress.stop()
        self.btn_convert.config(state="normal")
        self.btn_cancel.config(state="disabled")


if __name__ == "__main__":
    App().mainloop()